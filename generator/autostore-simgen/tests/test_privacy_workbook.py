"""Identifier/allowlist enforcement, pattern scans, workbook hygiene and overwrite protection."""

import re
import zipfile

import openpyxl
import pandas as pd
import pytest

from simgen.cli import main
from simgen.exporter import ExportBlocked, ExportError, verify_workbook
from simgen.identifiers import PATTERNS
from simgen.pipeline import RunOptions, build, export
from simgen.schema import SchemaDefinition
from simgen.security import DenylistMatcher, long_token_pattern, scan_text
from simgen.validate_file import validate_workbook_file
from simgen.validators.privacy import validate_supporting, validate_transactions
from simgen.security import build_vocabulary


def _small(tmp_path, **kw):
    opts = RunOptions(configs=["AMBIENT_BIN_MAINTENANCE"], target=300, seed=8, output=tmp_path / "amb_bm.xlsx", **kw)
    return build(opts)


def test_identifier_patterns(medium):
    kinds = {"from_tu": "tu", "to_tu": "tu", "from_product": "product", "to_product": "product",
             "from_lot": "lot", "to_lot": "lot", "from_order": "order", "transaction_id": "transaction",
             "from_lu": "lu", "to_lu": "lu", "port": "port_location"}
    for df in medium.frames.values():
        for col, kind in kinds.items():
            vals = df[col].dropna().astype(str).unique()
            assert all(PATTERNS[kind].match(v) for v in vals), col
        users = df["updated_by"].unique()
        assert all(u == "system" or PATTERNS["employee"].match(u) for u in users)
    all_tx = pd.concat([df["transaction_id"] for df in medium.frames.values()])
    assert all_tx.is_unique


def test_allowlist_injection_blocks_export(tmp_path):
    out = _small(tmp_path)
    df = out.frames["AMBIENT_BIN_MAINTENANCE"]
    df.loc[df.index[3], "from_product"] = "ACME-REAL-SKU-1234"  # not an approved identifier
    out.report.checks = [c for c in out.report.checks if c.category != "privacy"]
    out.report.extend(validate_transactions(out.config, "AMBIENT_BIN_MAINTENANCE", df, out.schema,
                                            build_vocabulary(out.config), DenylistMatcher([])))
    assert out.report.errors("privacy") > 0
    with pytest.raises(ExportBlocked):
        export(out)
    assert not (tmp_path / "amb_bm.xlsx").exists()
    assert not list(tmp_path.glob(".simgen-*"))


def test_blank_column_with_text_fails_allowlist(tmp_path):
    out = _small(tmp_path)
    df = out.frames["AMBIENT_BIN_MAINTENANCE"].copy()
    df.loc[df.index[0], "reference"] = "free text"
    checks = validate_transactions(out.config, "AMBIENT_BIN_MAINTENANCE", df, out.schema,
                                   build_vocabulary(out.config), DenylistMatcher([]))
    allow = [c for c in checks if c.name.startswith("Allowlist")][0]
    assert allow.violations == 1 and allow.failed


@pytest.mark.parametrize("value,kind", [
    ("contact jane.doe@example.com", "email address"),
    ("see https://intranet.example/x", "URL"),
    (r"C:\Users\someone\export.xlsx", "Windows path"),
    ("/home/someone/export.xlsx", "Unix path"),
    ("host 10.20.30.40", "IP address"),
    ("call 512-555-0199", "phone number"),
    ("key ghp_A1b2C3d4E5f6G7h8I9j0K1l2M3n4", "long token"),
])
def test_pattern_scan_catches_seeded_values(config, value, kind):
    assert kind in scan_text(value, long_token_pattern(24))
    checks = validate_supporting(config, {"Sheet": [[config.notice_text], ["note", value]]},
                                 {"creator": "Synthetic Data Generator", "lastModifiedBy": "Synthetic Data Generator"},
                                 DenylistMatcher([]), "file x")
    scan = [c for c in checks if "supporting sheets" in c.name][0]
    assert scan.failed


def test_vocabulary_does_not_trip_scans(medium, config):
    token = long_token_pattern(24)
    for df in medium.frames.values():
        for col in df.columns:
            if col.startswith("_") or df[col].dtype != object:
                continue
            for v in df[col].dropna().astype(str).unique()[:2000]:
                assert not scan_text(v, token), (col, "false positive")


def test_denylist_blocks(tmp_path):
    deny = tmp_path / "deny.txt"
    deny.write_text("# values that must never appear\nSIM_OWNER_01\n", encoding="utf-8")
    out = build(RunOptions(configs=["AMBIENT_BIN_MAINTENANCE"], target=200, seed=8, output=tmp_path / "x.xlsx",
                           denylist=deny))
    assert out.report.errors("privacy") > 0
    with pytest.raises(ExportBlocked):
        export(out)
    assert not (tmp_path / "x.xlsx").exists()


def test_workbook_hygiene_and_layout(tmp_path):
    out = _small(tmp_path)
    export(out)
    path = tmp_path / "amb_bm.xlsx"
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Transactions", "Synthetic Data Notice", "Validation Summary", "Employee Profiles",
                             "Schema Dictionary", "Calibration"]
    assert all(ws.sheet_state == "visible" for ws in wb.worksheets)
    ws = wb["Transactions"]
    assert ws.max_column == 89 and ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert [c.value for c in ws[1]] == SchemaDefinition().names
    assert ws.max_row == 301
    for name in wb.sheetnames[1:]:
        assert wb[name]["A1"].value == "SYNTHETIC TEST DATA - NOT FOR PRODUCTION USE"
    assert wb.properties.creator == "Synthetic Data Generator"
    assert wb.properties.lastModifiedBy == "Synthetic Data Generator"
    assert not wb.properties.title and not wb.properties.subject and not wb.properties.keywords
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        assert not any("vba" in n.lower() or "externallink" in n.lower() or "connections" in n.lower() for n in names)
        for n in names:
            if n.startswith("xl/worksheets/"):
                assert not re.search(r"<f[\s>]", z.read(n).decode("utf-8"))
        app = z.read("docProps/app.xml").decode("utf-8")
        assert "<Company>" not in app and "<Manager>" not in app
    checks = verify_workbook(path, out.config)
    assert not [c for c in checks if c.failed]
    # timestamps survive with millisecond precision
    back = pd.read_excel(path, sheet_name="Transactions")
    orig = out.frames["AMBIENT_BIN_MAINTENANCE"]["create_date"]
    assert (back["Create Date"].dt.round("ms").to_numpy() == orig.to_numpy().astype("datetime64[ms]")).all()


def test_round_trip_validation_of_written_file(tmp_path):
    out = build(RunOptions(configs=["COOLER_CYCLE_COUNT", "COOLER_BIN_MAINTENANCE"], target=400, seed=12,
                           output_dir=tmp_path, combined=True))
    export(out)
    for f in ("clr_cc.xlsx", "clr_bm.xlsx", "combined.xlsx"):
        report = validate_workbook_file(tmp_path / f)
        assert report.errors("logical") == 0, f
        assert report.errors("privacy") == 0, f
    wb = openpyxl.load_workbook(tmp_path / "combined.xlsx", read_only=True)
    assert wb.sheetnames[:2] == ["Transactions CLR_CC", "Transactions CLR_BM"]
    wb.close()
    assert (tmp_path / "validation_report.md").exists()


def test_overwrite_protection(tmp_path):
    out = _small(tmp_path)
    export(out)
    again = _small(tmp_path)
    with pytest.raises(ExportError, match="overwrite"):
        export(again)
    forced = _small(tmp_path, overwrite=True)
    export(forced)
    assert (tmp_path / "amb_bm.xlsx").exists()


def test_schema_file_override_and_cluster_id(tmp_path):
    names = [f"Field {i:02d}" for i in range(1, 90)]
    sf = tmp_path / "schema.yaml"
    sf.write_text("columns:\n" + "".join(f"  - \"{n}\"\n" for n in names), encoding="utf-8")
    out = build(RunOptions(configs=["AMBIENT_CYCLE_COUNT"], target=100, seed=2, output=tmp_path / "s.xlsx",
                           schema_file=sf, include_cluster_id=True))
    export(out)
    ws = openpyxl.load_workbook(tmp_path / "s.xlsx", read_only=True)["Transactions"]
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert header == names
    second = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert PATTERNS["cluster"].match(second[38])


def test_cli_generate_and_validate(tmp_path, capsys):
    rc = main(["generate", "--config", "CLR_BM", "--rows", "150", "--seed", "3", "--output", str(tmp_path / "c.xlsx"),
               "--quiet"])
    assert rc == 0 and (tmp_path / "c.xlsx").exists()
    rc = main(["validate", str(tmp_path / "c.xlsx")])
    assert rc == 0
    rc = main(["generate", "--config", "CLR_BM", "--rows", "1", "--output", str(tmp_path / "d.xlsx"), "--quiet"])
    assert rc == 1 and not (tmp_path / "d.xlsx").exists()
    rc = main(["generate", "--config", "CLR_BM", "--rows", "150", "--seed", "3", "--output", str(tmp_path / "c.xlsx"),
               "--quiet"])
    assert rc == 1  # refuses to overwrite
