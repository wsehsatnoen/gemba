"""Re-validate an exported workbook from the file alone.

Cluster keys are rebuilt from the rows (TU + Create Date order, a new cluster
starting at Cluster Sequence Number 1); the cluster employee is the human
updater when one is visible. Logical, privacy and package-hygiene checks run
exactly as at generation time, except checks that need generator internals.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .config import load_config
from .exporter import verify_workbook
from .schema import load_schema
from .security import DenylistMatcher, build_vocabulary, load_denylist
from .validators.logical import validate_config, validate_cross
from .validators.privacy import validate_transactions
from .validators.results import ValidationReport


def _read_sheet(path: Path, title: str) -> list[tuple]:
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb[title]
        return [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def frame_from_rows(rows: list[tuple], schema) -> pd.DataFrame:
    header = [str(h) for h in rows[0]]
    if header != schema.names:
        raise ValueError("Transactions header does not match the schema")
    data = list(rows[1:])
    df = pd.DataFrame(data, columns=schema.keys)
    for col in schema.columns:
        s = df[col.key]
        if col.dtype == "datetime":
            df[col.key] = pd.to_datetime(s).dt.round("ms")
        elif col.dtype == "date":
            df[col.key] = pd.Series([None if v is None or (isinstance(v, float) and np.isnan(v)) or v is pd.NaT
                                     else (v.date() if isinstance(v, dt.datetime) else v) for v in s], dtype=object)
        elif col.dtype == "int":
            df[col.key] = pd.array([None if v is None or pd.isna(v) else int(v) for v in s], dtype="Int64")
        else:
            df[col.key] = pd.Series([None if v is None or (isinstance(v, float) and np.isnan(v)) else v for v in s],
                                    dtype=object)
    return df


def reconstruct_internal(df: pd.DataFrame, system_user: str) -> pd.DataFrame:
    df = df.copy()
    order = df.sort_values(["from_tu", "create_date", "cluster_sequence_number"], kind="stable").index
    seq = df.loc[order, "cluster_sequence_number"].to_numpy()
    cluster = np.cumsum(seq == 1)
    df.loc[order, "_cluster"] = cluster
    df["_cluster"] = df["_cluster"].astype(int)
    humans = df["updated_by"].where(df["updated_by"] != system_user)
    emp = humans.groupby(df["_cluster"]).transform(lambda s: s.dropna().iloc[0] if s.notna().any() else None)
    df["_employee"] = emp.astype(object).where(emp.notna(), None)
    return df


def _config_for(config, df: pd.DataFrame) -> str:
    op = df["operation"].dropna().iloc[0]
    zone = df["zone"].dropna().iloc[0]
    for name, spec in config.configurations.items():
        if spec.operation == op and config.environments[spec.environment].zone == zone:
            return name
    raise ValueError("cannot determine the configuration of the Transactions sheet")


def validate_workbook_file(path: str | Path, config_file: str | Path | None = None,
                           denylist: str | Path | None = None, schema_file: str | Path | None = None) -> ValidationReport:
    path = Path(path)
    config = load_config(config_file)
    report = ValidationReport()
    deny = DenylistMatcher(load_denylist(denylist))
    wb = load_workbook(path, read_only=True)
    titles = wb.sheetnames
    wb.close()
    tx_titles = [t for t in titles if t.startswith("Transactions")]
    frames: dict[str, pd.DataFrame] = {}
    for title in tx_titles:
        rows = _read_sheet(path, title)
        schema = load_schema(schema_file)
        df = frame_from_rows(rows, schema)
        if df["cluster_id"].notna().any():
            schema = load_schema(schema_file, include_cluster_id=True)
        name = _config_for(config, df)
        df = reconstruct_internal(df, config.system_user)
        df["_config"] = name
        frames[name] = df
        report.extend(validate_config(config, name, df, schema, None))
        report.extend(validate_transactions(config, name, df, schema, build_vocabulary(config), deny))
    by_env: dict[str, dict[str, pd.DataFrame]] = {}
    for name, df in frames.items():
        by_env.setdefault(config.configurations[name].environment, {})[name] = df
    for env, fr in by_env.items():
        report.extend(validate_cross(config, fr, env))
    others = [t for t in titles if t not in tx_titles]
    report.extend(verify_workbook(path, config, expected_sheets=tx_titles + others, transaction_sheets=tx_titles,
                                  denylist=deny if deny else None))
    return report
