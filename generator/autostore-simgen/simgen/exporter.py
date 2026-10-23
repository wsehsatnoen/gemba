"""Workbook export with hygiene controls and atomic writes.

Supporting sheets are built with openpyxl. Transactions sheets are streamed as
SpreadsheetML directly (openpyxl alone needs ~30 s for 100,000 rows); the two
are merged by rewriting the saved package, then the result is verified and
atomically renamed into place.

Hygiene: creator / last-modified-by are "Synthetic Data Generator"; no
company, title or other personal properties; no macros, formulas, external
links, data connections or hidden sheets. The notice text is the first line of
every sheet except the Transactions sheets. Existing files are never
overwritten unless ``overwrite`` is set.
"""

from __future__ import annotations

import datetime as dt
import os
import re
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.sax.saxutils import escape

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .config import Config
from .schema import SchemaDefinition
from .security import DenylistMatcher, long_token_pattern, scan_text
from .validators.results import FAIL, PASS, CheckResult

EXCEL_EPOCH_DAYS = 25569  # 1970-01-01 as an Excel serial
DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss.000"
DATE_FORMAT = "yyyy-mm-dd"
INT_FORMAT = "0"
SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


class ExportError(RuntimeError):
    """Raised when a workbook cannot be written or fails verification."""


class ExportBlocked(ExportError):
    """Raised when logical or privacy failures block export."""


@dataclass
class SheetSpec:
    title: str
    rows: list[list[object]]
    widths: dict[int, float] | None = None
    header_row: int | None = None  # 1-based row index of a table header to bold


# ---------------------------------------------------------------------------------
# Streaming writer for Transactions sheets
# ---------------------------------------------------------------------------------

def _column_payloads(df: pd.DataFrame, schema: SchemaDefinition) -> list[tuple[str, object]]:
    """Per column: (kind, values) prepared for fast XML serialisation."""
    out: list[tuple[str, object]] = []
    for col in schema.columns:
        s = df[col.key]
        if col.dtype == "datetime":
            ms = s.to_numpy().astype("datetime64[ms]").astype(np.int64)
            serial = ms / 86_400_000.0 + EXCEL_EPOCH_DAYS
            out.append(("datetime", [repr(float(x)) for x in serial]))
        elif col.dtype == "date":
            vals = []
            base = dt.date(1899, 12, 30)
            for v in s.tolist():
                vals.append(None if v is None or (isinstance(v, float) and np.isnan(v)) or v is pd.NaT
                            else str((pd.Timestamp(v).date() - base).days))
            out.append(("date", vals))
        elif col.dtype == "int":
            vals = [None if v is pd.NA or v is None else str(int(v)) for v in s.tolist()]
            out.append(("int", vals))
        else:
            if s.isna().all():
                out.append(("empty", None))
            else:
                out.append(("string", [None if (v is None or v is pd.NA) else str(v) for v in s.tolist()]))
    return out


def stream_sheet_xml(fh, df: pd.DataFrame, schema: SchemaDefinition, styles: dict[str, int],
                     widths: list[float]) -> None:
    n_rows = len(df)
    n_cols = len(schema.columns)
    letters = [get_column_letter(i + 1) for i in range(n_cols)]
    last = f"{letters[-1]}{n_rows + 1}"
    write = fh.write
    write(b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n')
    write(f'<worksheet xmlns="{SHEET_NS}" xmlns:r="{REL_NS}">'.encode())
    write(f'<dimension ref="A1:{last}"/>'.encode())
    write(b'<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" '
          b'state="frozen"/><selection pane="bottomLeft" activeCell="A2" sqref="A2"/></sheetView></sheetViews>')
    write(b'<sheetFormatPr defaultRowHeight="15"/>')
    cols_xml = "".join(f'<col min="{i + 1}" max="{i + 1}" width="{w:.1f}" customWidth="1"/>' for i, w in enumerate(widths))
    write(f"<cols>{cols_xml}</cols>".encode())
    write(b"<sheetData>")
    hs = styles["header"]
    header = "".join(f'<c r="{letters[i]}1" t="inlineStr" s="{hs}"><is><t>{escape(c.name)}</t></is></c>'
                     for i, c in enumerate(schema.columns))
    write(f'<row r="1">{header}</row>'.encode())
    payloads = _column_payloads(df, schema)
    s_dt, s_date, s_int = styles["datetime"], styles["date"], styles["int"]
    str_cache: dict[str, str] = {}
    active = [(i, kind, vals) for i, (kind, vals) in enumerate(payloads) if kind != "empty"]
    chunk: list[str] = []
    for r in range(n_rows):
        rn = r + 2
        parts = [f'<row r="{rn}">']
        for i, kind, vals in active:
            v = vals[r]  # type: ignore[index]
            if v is None:
                continue
            ref = f"{letters[i]}{rn}"
            if kind == "string":
                frag = str_cache.get(v)
                if frag is None:
                    frag = f' t="inlineStr"><is><t>{escape(v)}</t></is></c>'
                    str_cache[v] = frag
                parts.append(f'<c r="{ref}"{frag}')
            elif kind == "datetime":
                parts.append(f'<c r="{ref}" s="{s_dt}"><v>{v}</v></c>')
            elif kind == "date":
                parts.append(f'<c r="{ref}" s="{s_date}"><v>{v}</v></c>')
            else:
                parts.append(f'<c r="{ref}" s="{s_int}"><v>{v}</v></c>')
        parts.append("</row>")
        chunk.append("".join(parts))
        if len(chunk) >= 2000:
            write("".join(chunk).encode("utf-8"))
            chunk = []
    if chunk:
        write("".join(chunk).encode("utf-8"))
    write(b"</sheetData>")
    write(f'<autoFilter ref="A1:{last}"/>'.encode())
    write(b'<pageMargins left="0.75" right="0.75" top="1" bottom="1" header="0.5" footer="0.5"/>')
    write(b"</worksheet>")


# ---------------------------------------------------------------------------------
# Package helpers
# ---------------------------------------------------------------------------------

def _core_xml(creator: str, when: dt.datetime) -> bytes:
    stamp = when.strftime("%Y-%m-%dT%H:%M:%SZ")
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
        'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        f"<dc:creator>{escape(creator)}</dc:creator><cp:lastModifiedBy>{escape(creator)}</cp:lastModifiedBy>"
        f'<dcterms:created xsi:type="dcterms:W3CDTF">{stamp}</dcterms:created>'
        f'<dcterms:modified xsi:type="dcterms:W3CDTF">{stamp}</dcterms:modified>'
        "</cp:coreProperties>"
    ).encode("utf-8")


def _app_xml(creator: str) -> bytes:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        f"<Application>{escape(creator)}</Application><DocSecurity>0</DocSecurity><ScaleCrop>false</ScaleCrop>"
        "<LinksUpToDate>false</LinksUpToDate><SharedDoc>false</SharedDoc><HyperlinksChanged>false</HyperlinksChanged>"
        "</Properties>"
    ).encode("utf-8")


def _sheet_parts(z: zipfile.ZipFile) -> dict[str, str]:
    """Map sheet title -> part name using workbook.xml and its relationships."""
    wb = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rel_target = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels)}
    rel_target.update({m.group(2): m.group(1) for m in re.finditer(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
    out: dict[str, str] = {}
    for m in re.finditer(r"<sheet\b([^>]*)/>", wb):
        attrs = m.group(1)
        name = re.search(r'name="([^"]+)"', attrs)
        rid = re.search(r'r:id="([^"]+)"', attrs)
        if not name or not rid:
            continue
        target = rel_target[rid.group(1)]
        target = target.lstrip("/")
        if not target.startswith("xl/"):
            target = "xl/" + target
        out[_unescape(name.group(1))] = target
    return out


def _unescape(text: str) -> str:
    return text.replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"').replace("&apos;", "'").replace("&amp;", "&")


# ---------------------------------------------------------------------------------
# Writing
# ---------------------------------------------------------------------------------

def _transaction_widths(schema: SchemaDefinition) -> list[float]:
    widths = []
    for c in schema.columns:
        if c.dtype == "datetime":
            widths.append(24.0)
        elif c.dtype == "date":
            widths.append(13.0)
        else:
            widths.append(max(10.0, min(30.0, len(c.name) + 4.0)))
    return widths


def write_workbook(path: str | Path, transactions: dict[str, pd.DataFrame], supporting: list[SheetSpec],
                   schema: SchemaDefinition, config: Config, overwrite: bool = False,
                   generated_at: dt.datetime | None = None,
                   denylist: DenylistMatcher | None = None) -> list[CheckResult]:
    """Write a workbook atomically and return its post-write verification checks."""
    path = Path(path)
    if path.exists() and not overwrite:
        raise ExportError(f"refusing to overwrite existing file '{path.name}' (pass --overwrite)")
    path.parent.mkdir(parents=True, exist_ok=True)
    creator = str(config.raw["meta"]["workbook_creator"])
    when = generated_at or dt.datetime.now(dt.timezone.utc).replace(tzinfo=None)

    wb = Workbook()
    wb.remove(wb.active)
    styles: dict[str, int] = {}
    for title in transactions:
        ws = wb.create_sheet(title)
        cell = ws.cell(row=1, column=1, value="placeholder")
        cell.font = Font(bold=True)
        styles["header"] = cell.style_id
        c_dt = ws.cell(row=2, column=1, value=dt.datetime(2027, 1, 1))
        c_dt.number_format = DATETIME_FORMAT
        styles["datetime"] = c_dt.style_id
        c_d = ws.cell(row=2, column=2, value=dt.date(2027, 1, 1))
        c_d.number_format = DATE_FORMAT
        styles["date"] = c_d.style_id
        c_i = ws.cell(row=2, column=3, value=1)
        c_i.number_format = INT_FORMAT
        styles["int"] = c_i.style_id
    for spec in supporting:
        ws = wb.create_sheet(spec.title)
        for row in spec.rows:
            ws.append([_cell_value(v) for v in row])
        ws["A1"].font = Font(bold=True)
        if spec.header_row:
            for c in ws[spec.header_row]:
                c.font = Font(bold=True)
        for idx, width in (spec.widths or {}).items():
            ws.column_dimensions[get_column_letter(idx)].width = width
        for r in ws.iter_rows():
            for c in r:
                if isinstance(c.value, dt.datetime):
                    c.number_format = "yyyy-mm-dd hh:mm:ss"
                elif isinstance(c.value, dt.date):
                    c.number_format = DATE_FORMAT
    wb.active = 0
    wb.properties.creator = creator
    wb.properties.lastModifiedBy = creator

    fd1, tmp1 = tempfile.mkstemp(prefix=".simgen-", suffix=".stage.xlsx", dir=str(path.parent))
    os.close(fd1)
    fd2, tmp2 = tempfile.mkstemp(prefix=".simgen-", suffix=".tmp.xlsx", dir=str(path.parent))
    os.close(fd2)
    try:
        wb.save(tmp1)
        widths = _transaction_widths(schema)
        with zipfile.ZipFile(tmp1, "r") as zin:
            parts = _sheet_parts(zin)
            replace = {parts[title]: title for title in transactions}
            with zipfile.ZipFile(tmp2, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as zout:
                for item in zin.infolist():
                    name = item.filename
                    if name in replace:
                        with zout.open(name, "w", force_zip64=True) as fh:
                            stream_sheet_xml(fh, transactions[replace[name]], schema, styles, widths)
                    elif name == "docProps/core.xml":
                        zout.writestr(name, _core_xml(creator, when))
                    elif name == "docProps/app.xml":
                        zout.writestr(name, _app_xml(creator))
                    else:
                        zout.writestr(item, zin.read(name))
        checks = verify_workbook(tmp2, config, expected_sheets=list(transactions) + [s.title for s in supporting],
                                 transaction_sheets=list(transactions), denylist=denylist)
        if any(c.status == FAIL for c in checks):
            raise ExportBlocked("written workbook failed verification: " +
                                "; ".join(f"{c.name} ({c.violations})" for c in checks if c.status == FAIL))
        if path.exists() and not overwrite:
            raise ExportError(f"refusing to overwrite existing file '{path.name}' (pass --overwrite)")
        umask = os.umask(0)
        os.umask(umask)
        os.chmod(tmp2, 0o666 & ~umask)
        os.replace(tmp2, path)
        return checks
    finally:
        for tmp in (tmp1, tmp2):
            try:
                os.unlink(tmp)
            except FileNotFoundError:
                pass


def _cell_value(v: object) -> object:
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating,)):
        return float(v)
    if isinstance(v, (np.bool_, bool)):
        return "yes" if v else "no"
    return v


# ---------------------------------------------------------------------------------
# Verification
# ---------------------------------------------------------------------------------

_T_TEXT = re.compile(r"<t(?:\s[^>]*)?>([^<]*)</t>")


def verify_workbook(path: str | Path, config: Config, expected_sheets: list[str] | None = None,
                    transaction_sheets: list[str] | None = None,
                    denylist: DenylistMatcher | None = None) -> list[CheckResult]:
    """Package-level checks on a written workbook (aggregate results only)."""
    creator = str(config.raw["meta"]["workbook_creator"])
    notice = config.notice_text
    token_re = long_token_pattern(int(config.raw["validation"]["pattern_scan"]["long_token_min_length"]))
    checks: list[CheckResult] = []

    def add(category: str, name: str, violations: int, detail: str = "") -> None:
        checks.append(CheckResult(category, "workbook", name, PASS if violations == 0 else FAIL, violations, detail))

    with zipfile.ZipFile(path, "r") as z:
        names = z.namelist()
        lower = [n.lower() for n in names]
        macros = sum(1 for n in lower if n.endswith("vbaproject.bin") or n.endswith(".bin") and "vba" in n)
        content_types = z.read("[Content_Types].xml").decode("utf-8")
        macros += content_types.count("macroEnabled")
        add("logical", "No macros", macros)
        ext = sum(1 for n in lower if n.startswith("xl/externallinks/"))
        add("logical", "No external links", ext)
        conns = sum(1 for n in lower if n in ("xl/connections.xml",) or n.startswith("xl/querytables/"))
        add("logical", "No data connections", conns)
        wb_xml = z.read("xl/workbook.xml").decode("utf-8")
        hidden = len(re.findall(r'<sheet\b[^>]*state="(?:hidden|veryHidden)"', wb_xml))
        add("logical", "No hidden sheets", hidden)
        sheet_map = _sheet_parts(z)
        titles = list(sheet_map)
        if expected_sheets is not None:
            add("logical", "Sheet layout as specified", 0 if titles == expected_sheets else 1,
                f"{len(titles)} sheets")
        formulas = 0
        pattern_hits = 0
        deny_hits = 0
        notice_missing = 0
        if transaction_sheets is None:
            transaction_sheets = [t for t in titles if t.startswith("Transactions")]
        trans = set(transaction_sheets)
        col_issue = 0
        for title, part in sheet_map.items():
            xml = z.read(part).decode("utf-8")
            formulas += len(re.findall(r"<f[\s>]", xml))
            texts = [_unescape(t) for t in _T_TEXT.findall(xml)]
            if title not in trans:
                first = re.search(r'<c r="A1"[^>]*>(.*?)</c>', xml)
                first_text = ""
                if first:
                    m = re.search(r"<t(?:\s[^>]*)?>([^<]*)</t>", first.group(1))
                    if m:
                        first_text = _unescape(m.group(1))
                    else:
                        v = re.search(r"<v>(\d+)</v>", first.group(1))
                        if v and "sharedStrings" in "".join(names):
                            first_text = _shared_string(z, int(v.group(1)))
                if first_text != notice:
                    notice_missing += 1
            else:
                header = re.search(r'<row r="1">(.*?)</row>', xml)
                n_cols = len(re.findall(r"<c ", header.group(1))) if header else 0
                if n_cols != 89:
                    col_issue += 1
            for t in set(texts):
                if scan_text(t, token_re):
                    pattern_hits += 1
                if denylist and denylist.matches(t):
                    deny_hits += 1
        if "xl/sharedStrings.xml" in names:
            for t in set(_T_TEXT.findall(z.read("xl/sharedStrings.xml").decode("utf-8"))):
                t = _unescape(t)
                if scan_text(t, token_re):
                    pattern_hits += 1
                if denylist and denylist.matches(t):
                    deny_hits += 1
        add("logical", "No formulas", formulas)
        add("logical", "Transactions sheets have exactly 89 columns", col_issue)
        add("logical", "Notice text is the first line of every non-Transactions sheet", notice_missing)
        core = z.read("docProps/core.xml").decode("utf-8") if "docProps/core.xml" in names else ""
        app = z.read("docProps/app.xml").decode("utf-8") if "docProps/app.xml" in names else ""
        meta_bad = 0
        if f"<dc:creator>{escape(creator)}</dc:creator>" not in core:
            meta_bad += 1
        if f"<cp:lastModifiedBy>{escape(creator)}</cp:lastModifiedBy>" not in core:
            meta_bad += 1
        for tag in ("dc:title", "dc:subject", "dc:description", "cp:keywords", "cp:category", "dc:identifier"):
            if f"<{tag}>" in core:
                meta_bad += 1
        for tag in ("Company", "Manager", "HyperlinkBase"):
            if f"<{tag}>" in app:
                meta_bad += 1
        meta_bad += 1 if "docProps/custom.xml" in names else 0
        add("privacy", "Workbook metadata (creator, last-modified-by, no personal properties)", meta_bad)
        meta_text = " ".join(_T_TEXT.findall(core)) + re.sub(r"<[^>]+>", " ", core + " " + app)
        add("privacy", "Pattern scan of all sheets and metadata (written file)",
            pattern_hits + len(scan_text(meta_text, token_re)))
        if denylist:
            meta_deny = 1 if denylist.matches(meta_text) else 0
            add("privacy", "Deny-list scan of the written file", deny_hits + meta_deny, f"{deny_hits + meta_deny} match(es)")
    return checks


def _shared_string(z: zipfile.ZipFile, index: int) -> str:
    xml = z.read("xl/sharedStrings.xml").decode("utf-8")
    items = re.findall(r"<si>(.*?)</si>", xml, flags=re.S)
    if index >= len(items):
        return ""
    return _unescape("".join(re.findall(r"<t(?:\s[^>]*)?>([^<]*)</t>", items[index])))
